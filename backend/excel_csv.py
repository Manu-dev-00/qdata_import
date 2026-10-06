import io
import csv
from datetime import datetime
from typing import Any
import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, numbers
)
from openpyxl.utils import get_column_letter

# ============================================================
#  MODULE 3 — GÉNÉRATEUR EXCEL + CONVERTISSEUR CSV
#  Excel professionnel avec cellules suspectes en rouge
#  CSV prêt pour le job Qdata (station par station)
# ============================================================

# Colonnes Qdata standard — à adapter selon configuration
QDATA_COLUMNS = [
    "DATE",
    "TN",       # Température minimale
    "TX",       # Température maximale
    "TMOY",     # Température moyenne
    "RR 06:00", # Précipitations
    "EVP 06:00",# Evaporation
    "HR",       # Humidité relative
    "INSOLATION",
    # Ajouter ici les paramètres spécifiques à vos stations
]

# Styles Excel
HDR_BG    = "1B3A6B"
HDR_FG    = "FFFFFF"
META_BG   = "E8F0FE"
ALT_BG    = "F5F7FA"
SUSPECT_BG= "FFE0E0"
SUSPECT_FG= "CC0000"
OK_BG     = "E8F5E9"
ZERO_BG   = "FFF9E6"


def _thin_border():
    s = Side(style="thin", color="CCCCCC")
    return Border(left=s, right=s, top=s, bottom=s)


def _header_fill():
    return PatternFill("solid", fgColor=HDR_BG)


def generate_excel(
    station_id: str,
    station_nom: str,
    mois: str,
    annee: int,
    rows: list[dict],
    suspects_map: dict[str, list[str]],
    columns: list[str] = None,
) -> bytes:
    """
    Génère un fichier Excel formaté pour une station.
    - En-tête colorée avec infos station
    - Cellules suspectes en rouge (à vérifier manuellement)
    - Cellules avec 0 converti en jaune pâle
    - Ligne TOTAL et MOYENNE automatiques
    - Feuille de résumé anomalies
    """
    cols = columns or QDATA_COLUMNS
    wb = openpyxl.Workbook()

    # ── Feuille principale : données ──────────────────────────
    ws = wb.active
    ws.title = f"{station_id}_{mois[:3]}_{annee}"

    border = _thin_border()

    # Titre station
    ws.merge_cells(f"A1:{get_column_letter(len(cols))}1")
    ws["A1"] = f"Station : {station_id} — {station_nom}"
    ws["A1"].font = Font(bold=True, size=14, color=HDR_FG, name="Arial")
    ws["A1"].fill = PatternFill("solid", fgColor=HDR_BG)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    # Infos période
    ws.merge_cells(f"A2:{get_column_letter(len(cols))}2")
    ws["A2"] = (
        f"Période : {mois} {annee}  |  "
        f"Généré le {datetime.now().strftime('%d/%m/%Y %H:%M')}  |  "
        f"Vérifier les cellules rouges avant conversion CSV"
    )
    ws["A2"].font = Font(italic=True, size=10, color="444444", name="Arial")
    ws["A2"].fill = PatternFill("solid", fgColor=META_BG)
    ws["A2"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[2].height = 18

    # Légende
    ws.merge_cells(f"A3:{get_column_letter(len(cols))}3")
    ws["A3"] = (
        "Rouge = valeur suspecte à vérifier  |  "
        "Jaune = 0 converti automatiquement  |  "
        "Vert = valeur numérique valide"
    )
    ws["A3"].font = Font(size=9, color="666666", name="Arial")
    ws["A3"].fill = PatternFill("solid", fgColor="F0F0F0")
    ws["A3"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[3].height = 16

    # En-têtes colonnes
    for ci, col in enumerate(cols, 1):
        cell = ws.cell(row=4, column=ci, value=col)
        cell.font = Font(bold=True, size=11, color=HDR_FG, name="Arial")
        cell.fill = _header_fill()
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
        ws.column_dimensions[get_column_letter(ci)].width = max(len(col) + 4, 12)
    ws.row_dimensions[4].height = 22

    # Données
    for ri, row in enumerate(rows, 5):
        date_val = row.get("DATE", "")
        suspect_fields = suspects_map.get(str(date_val), [])
        is_alt = (ri % 2 == 0)

        for ci, col in enumerate(cols, 1):
            val = row.get(col, 0)
            cell = ws.cell(row=ri, column=ci, value=val)
            cell.font = Font(size=10, name="Arial")
            cell.border = border
            cell.alignment = Alignment(horizontal="center")

            if col == "DATE":
                cell.font = Font(bold=True, size=10, name="Arial")
                cell.alignment = Alignment(horizontal="left")
                continue

            # Coloration selon statut
            if col in suspect_fields:
                # Rouge : valeur originale suspecte
                cell.fill = PatternFill("solid", fgColor=SUSPECT_BG)
                cell.font = Font(size=10, name="Arial", color=SUSPECT_FG, bold=True)
            elif val == 0:
                # Jaune pâle : 0 converti
                cell.fill = PatternFill("solid", fgColor=ZERO_BG)
            else:
                # Vert : valeur numérique valide
                cell.fill = PatternFill("solid", fgColor=OK_BG if not is_alt else "F0FAF0")

    data_start_row = 5
    data_end_row   = 4 + len(rows)

    # Ligne TOTAL
    total_row = data_end_row + 1
    ws.cell(row=total_row, column=1, value="TOTAL").font = Font(
        bold=True, size=10, name="Arial"
    )
    ws.cell(row=total_row, column=1).fill = PatternFill("solid", fgColor="DDE3ED")
    for ci, col in enumerate(cols[1:], 2):
        col_letter = get_column_letter(ci)
        cell = ws.cell(
            row=total_row, column=ci,
            value=f"=SUM({col_letter}{data_start_row}:{col_letter}{data_end_row})"
        )
        cell.font = Font(bold=True, size=10, name="Arial")
        cell.fill = PatternFill("solid", fgColor="DDE3ED")
        cell.border = border
        cell.alignment = Alignment(horizontal="center")

    # Ligne MOYENNE
    mean_row = total_row + 1
    ws.cell(row=mean_row, column=1, value="MOYENNE").font = Font(
        bold=True, size=10, name="Arial"
    )
    ws.cell(row=mean_row, column=1).fill = PatternFill("solid", fgColor="DDE3ED")
    for ci, col in enumerate(cols[1:], 2):
        col_letter = get_column_letter(ci)
        cell = ws.cell(
            row=mean_row, column=ci,
            value=f"=AVERAGE({col_letter}{data_start_row}:{col_letter}{data_end_row})"
        )
        cell.font = Font(bold=True, size=10, name="Arial")
        cell.fill = PatternFill("solid", fgColor="DDE3ED")
        cell.border = border
        cell.alignment = Alignment(horizontal="center")
        cell.number_format = "0.00"

    # Figer les 4 premières lignes + colonne DATE
    ws.freeze_panes = "B5"

    # ── Feuille anomalies ──────────────────────────────────────
    ws_anom = wb.create_sheet(title="Anomalies_a_verifier")
    ws_anom["A1"] = "ANOMALIES DÉTECTÉES — À CORRIGER AVANT IMPORT QDATA"
    ws_anom["A1"].font = Font(bold=True, size=13, color="CC0000", name="Arial")
    ws_anom.merge_cells("A1:E1")

    anom_hdrs = ["Date", "Champs suspects", "Valeurs originales", "Action requise", "Corrigé"]
    for ci, h in enumerate(anom_hdrs, 1):
        cell = ws_anom.cell(row=2, column=ci, value=h)
        cell.font = Font(bold=True, color=HDR_FG, name="Arial")
        cell.fill = _header_fill()
        cell.border = border
        cell.alignment = Alignment(horizontal="center")
        ws_anom.column_dimensions[get_column_letter(ci)].width = 24

    anom_row = 3
    has_anomalies = False
    for date_str, fields in suspects_map.items():
        if fields:
            has_anomalies = True
            ws_anom.cell(row=anom_row, column=1, value=date_str)
            ws_anom.cell(row=anom_row, column=2, value=", ".join(fields))
            ws_anom.cell(row=anom_row, column=3, value="Voir feuille principale")
            ws_anom.cell(row=anom_row, column=4, value="Vérifier sur fiche papier et corriger")
            chk = ws_anom.cell(row=anom_row, column=5, value="Non")
            chk.font = Font(color="CC0000", name="Arial")
            for ci in range(1, 6):
                ws_anom.cell(row=anom_row, column=ci).border = border
                ws_anom.cell(row=anom_row, column=ci).fill = PatternFill("solid", fgColor=SUSPECT_BG)
            anom_row += 1

    if not has_anomalies:
        ws_anom["A3"] = "Aucune anomalie detectee — donnees propres"
        ws_anom["A3"].font = Font(color="059669", bold=True, name="Arial")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def excel_to_csv(
    rows: list[dict],
    station_id: str,
    columns: list[str] = None,
    separator: str = ";",
) -> str:
    """
    Convertit les données vérifiées en CSV prêt pour le job Qdata.
    Format station par station (pas besoin d'ID station dans chaque ligne).
    Première colonne = DATE, puis les paramètres dans l'ordre Qdata.
    """
    cols = columns or QDATA_COLUMNS
    output = io.StringIO()
    writer = csv.DictWriter(
        output,
        fieldnames=cols,
        delimiter=separator,
        extrasaction="ignore",
        lineterminator="\r\n",
    )
    writer.writeheader()
    for row in rows:
        clean_row = {col: row.get(col, 0) for col in cols}
        writer.writerow(clean_row)
    return output.getvalue()


def generate_batch_report(results: list[dict]) -> bytes:
    """
    Génère un rapport Excel récapitulatif de tous les imports
    (utile pour suivre les 240 stations).
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rapport_global"
    border = _thin_border()

    ws.merge_cells("A1:G1")
    ws["A1"] = f"RAPPORT GLOBAL D'IMPORT QDATA — {datetime.now().strftime('%d/%m/%Y')}"
    ws["A1"].font = Font(bold=True, size=14, color=HDR_FG, name="Arial")
    ws["A1"].fill = _header_fill()
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 26

    hdrs = ["ID Station", "Nom", "Mois", "Année", "Statut", "Anomalies", "CSV généré"]
    for ci, h in enumerate(hdrs, 1):
        cell = ws.cell(row=2, column=ci, value=h)
        cell.font = Font(bold=True, color=HDR_FG, name="Arial")
        cell.fill = _header_fill()
        cell.border = border
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(ci)].width = 18

    for ri, res in enumerate(results, 3):
        status = res.get("statut", "en attente")
        is_alt = ri % 2 == 0
        row_data = [
            res.get("station_id", ""),
            res.get("station_nom", ""),
            res.get("mois", ""),
            res.get("annee", ""),
            status,
            res.get("anomalies", 0),
            "Oui" if res.get("csv_genere") else "Non",
        ]
        for ci, val in enumerate(row_data, 1):
            cell = ws.cell(row=ri, column=ci, value=val)
            cell.font = Font(size=10, name="Arial")
            cell.border = border
            cell.alignment = Alignment(horizontal="center")
            if status == "OK":
                cell.fill = PatternFill("solid", fgColor="E8F5E9" if not is_alt else "F0FAF0")
            elif status == "anomalies":
                cell.fill = PatternFill("solid", fgColor=ZERO_BG)
            else:
                cell.fill = PatternFill("solid", fgColor="F5F5F5")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
